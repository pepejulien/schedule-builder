#!/usr/bin/env python3
"""Drivers-page preferences in the solver + the rules engine (2026-10).

Config keys (all optional; absent = the old behavior exactly):
  extra_unavailable {name: {Day: why}}  HARD day off (like a submitted Unavailable)
  driver_max_days   {name: n}           HARD cap on days on the clock this week
  driver_max_row    {name: n}           HARD max days in a row (last week counts)
  keep_together     [name]              SOFT: days next to each other
  prefer_days       {name: [Day]}       SOFT: placement boost
Names are matched loosely; an unmatched / ambiguous one is skipped and listed.
In runner.py a manual edit past a driver's own cap is POLICY (allowed, flagged)
and a Drivers-page day off is the typed-confirm 'unavail' bucket.

Run:  python tests/gen_fixtures.py   (once)  then  python tests/driver_prefs_solver_test.py
"""
import contextlib
import io
import json
import os
import random
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, '..', 'public', 'solver'))
sys.path.insert(0, HERE)
import openpyxl  # noqa: E402
import build_weekly_schedule as B  # noqa: E402  (the CURRENT solver, public/solver)
import runner  # noqa: E402

ALL = B.ALL
FIX = os.path.join(HERE, 'fixtures')
ONE = __import__('datetime').timedelta(days=1)
fails = []
tmp = tempfile.mkdtemp()


def ok(cond, msg):
    print(('  ok   ' if cond else '  FAIL ') + msg)
    if not cond:
        fails.append(msg)


def J(fn, payload=None):
    return json.loads(fn(json.dumps(payload or {})))


def fixture_cfg(**extra):
    cfg = json.load(open(os.path.join(FIX, 'Week-40-config.json'), encoding='utf-8'))
    for k in ('prev_week_file', 'prefs_csv', 'avail_file'):
        cfg[k] = os.path.join(FIX, os.path.basename(cfg[k]))
    cfg['out'] = os.path.join(tmp, 'fx.xlsx')
    cfg.update(extra)
    return cfg


def save_cfg(cfg, name='c.json'):
    p = os.path.join(tmp, name)
    json.dump(cfg, open(p, 'w', encoding='utf-8'))
    return p


def build(cfg):
    res = B.build_schedule(B.load_config(save_cfg(cfg)))
    return res, B.check_invariants(res)


def synth_week(names, per_day, tag, unav=None, prev=None, **extra):
    """A synthetic week: every driver available unless `unav` says otherwise;
    `prev` = {name: [days worked LAST week]} for the days-in-a-row carryover."""
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Shifts & Availability'
    ws.append(['Associate Name', 'Transporter ID'] + ALL)
    for i, n in enumerate(names):
        ws.append([n, f'T{i:04d}'] + ['Unavailable' if d in (unav or {}).get(n, ()) else ''
                                      for d in ALL])
    avail = os.path.join(tmp, f'{tag}-avail.xlsx'); wb.save(avail)
    pf = None
    if prev:
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Shifts & Availability'
        ws.append(['Associate Name', 'Transporter ID'] + ALL)
        for i, n in enumerate(names):
            ws.append([n, f'T{i:04d}'] + ['10:45 AM' if d in prev.get(n, ()) else '' for d in ALL])
        pf = os.path.join(tmp, f'{tag}-prev.xlsx'); wb.save(pf)
    cfg = dict(start_date='2026-08-02', waves={d: {'10:25 AM': per_day} for d in ALL},
               backup_per_day={d: 0 for d in ALL}, max_primary_days=4, weekly_hours_cap=40,
               max_total_days=5, free_total_days=4, prev_week_file=pf, avail_file=avail,
               out=os.path.join(tmp, f'{tag}-out.xlsx'), weekend_spread=False)
    cfg.update(extra)
    return cfg


def days_on(dr):
    return set(dr['prim']) | set(dr['bk']) | set(dr['helper'])


def total(dr):
    return (len(dr['prim']) + len(dr['helper']) + len(dr['bk'])
            + len(dr['extra']) + len(dr['meet']))


def drv(res, name):
    return next(dr for dr in res.roster if dr['name'] == name)


def longest_run_touching_week(res, dr):
    s = (set(dr['w_prev']) | {res.DATEALL[d] for d in days_on(dr)}
         | {res.DATEALL[d] for d in dr['extra'] | dr['meet']})
    lo, hi, best = res.DATEALL['Sun'], res.DATEALL['Sat'], 0
    for x in s:
        if x - ONE in s:
            continue
        n = 1
        while x + n * ONE in s:
            n += 1
        if x <= hi and x + (n - 1) * ONE >= lo:
            best = max(best, n)
    return best


HARD = ('PREF-', 'UNAVAIL', 'CONSEC', 'DAYS7', 'TOTDAYS', 'DAYCAP', 'OT:', 'DUP-DAY', 'TARGET')


def hard_errs(chk):
    return [e for e in chk['errors'] if e.startswith(HARD)]


# =========================================================== A. the solver ====
print('A1 extra_unavailable: a hard day off, written as Unavailable')
EU = {'aaron bell': {'Tue': 'timeoff', 'Wed': 'days', 'Thu': 'days'},
      'Colin Drake': {'Sat': 'weekend', 'Tue': 'timeoff'}}     # Colin's Tue is a seed day
res, chk = build(fixture_cfg(extra_unavailable=EU))
a, c = drv(res, 'Aaron Bell'), drv(res, 'Colin Drake')
ok(a['unav_why'] == {'Tue': 'timeoff', 'Wed': 'days', 'Thu': 'days'}, f"unav_why kept: {a['unav_why']}")
ok(c['unav_why'] == {'Sat': 'weekend', 'Tue': 'timeoff'}, f"Colin unav_why: {c['unav_why']}")
ok(not (days_on(a) & {'Tue', 'Wed', 'Thu'}), f"Aaron not on Tue/Wed/Thu: {sorted(days_on(a))}")
ok(not (days_on(c) & {'Sat', 'Tue'}), f"Colin not on Sat/Tue (seed Tue dropped): {sorted(days_on(c))}")
ok(not hard_errs(chk), f'no hard-rule errors: {hard_errs(chk)}')
with contextlib.redirect_stdout(io.StringIO()):
    B.write_xlsx(res)
ws = openpyxl.load_workbook(res.cfg['out'])['Shifts & Availability']
hdr = {str(ws.cell(4, col).value or '')[:3]: col for col in range(1, ws.max_column + 1)}
rows = {ws.cell(r, 1).value: r for r in range(6, ws.max_row + 1)}
ok(all(ws.cell(rows['Aaron Bell'], hdr[d]).value == 'Unavailable' for d in ('Tue', 'Wed', 'Thu')),
   'workbook: Aaron Tue/Wed/Thu read "Unavailable" like a submitted day')
ok(ws.cell(rows['Colin Drake'], hdr['Sat']).value == 'Unavailable', 'workbook: Colin Sat "Unavailable"')
buf = io.StringIO()
with contextlib.redirect_stdout(buf):
    B.print_summary(res, chk)
ok('Drivers-page days off added: 5 day(s) for 2 driver(s)' in buf.getvalue(),
   'summary prints the days-off line')

print('A2 driver_max_days: a hard cap, even on a Top who would get 5 / an exact target')
TOPS = [f'Topz{i:02d}' for i in range(8)]
FAIRS = [f'Fairz{i:02d}' for i in range(8)]
base = synth_week(TOPS + FAIRS + ['Exactz Ed'], 8, 'cap', most_days=TOPS,
                  exact_days={'Exactz Ed': 4}, backup_per_day={d: 3 for d in ALL})
res0, chk0 = build(base)
ok(total(drv(res0, 'Topz00')) == 5, f"precondition: Topz00 gets 5 days uncapped ({total(drv(res0, 'Topz00'))})")
ok(len(drv(res0, 'Exactz Ed')['prim']) == 4, 'precondition: Exactz Ed gets its exact 4')
capcfg = dict(base, driver_max_days={'Topz00': 3, 'Topz01': 2, 'Exactz Ed': 2, 'Fairz00': 1,
                                     'Fairz01': 0})
res, chk = build(capcfg)
for n, k in (('Topz00', 3), ('Topz01', 2), ('Exactz Ed', 2), ('Fairz00', 1), ('Fairz01', 0)):
    ok(total(drv(res, n)) <= k, f'{n} capped at {k}: has {total(drv(res, n))}')
ok(total(drv(res, 'Topz00')) == 3, 'Topz00 still gets the 3 it allows')
ok(('exactz ed', 4, 2) in chk['target_short'], f"exact target above the cap reported short: {chk['target_short']}")
ok(not chk['target_bad'], 'not a TARGET error')
ok(not hard_errs(chk), f'no hard-rule errors: {hard_errs(chk)}')
ok(any(n == 'Topz00' and k == 3 and t == 3 for n, k, t in chk['pref_caps']), 'pref_caps lists Topz00 at cap')
# a meeting / dispatch day counts toward the cap
capcfg2 = dict(base, driver_max_days={'Topz02': 2}, extra_worked_days={'Topz02': ['Mon']})
res, chk = build(capcfg2)
ok(total(drv(res, 'Topz02')) <= 2 and len(drv(res, 'Topz02')['prim']) <= 1,
   f"dispatch day counts toward max_days: {total(drv(res, 'Topz02'))} days")
# a training pair that only the cap blocks -> reported, not a build error
res, chk = build(dict(base, training_pairs=[{'trainer': 'Topz03', 'trainee': 'Fairz05'}],
                      driver_max_days={'Topz03': 0}))
ok(not res.PAIRLOG and not any('TRAINING' in x for x in res.infeasible)
   and any('TRAINING' in x and 'Drivers-page' in x for x in res.notes),
   'training pair blocked by a cap -> a note, not infeasible')

print('A3 driver_max_row: max days in a row, last week counts')
ROWS = [f'Rowk{i:02d}' for i in range(12)]
cfg = synth_week(ROWS, 7, 'row', most_days=ROWS[:4],
                 prev={'Rowk00': ['Sat'], 'Rowk01': ['Fri', 'Sat'], 'Rowk02': ['Sat']},
                 driver_max_row={'Rowk00': 1, 'Rowk01': 2, 'Rowk02': 2, 'Rowk03': 1})
res, chk = build(cfg)
for n, k in (('Rowk00', 1), ('Rowk01', 2), ('Rowk02', 2), ('Rowk03', 1)):
    dr = drv(res, n)
    ok(longest_run_touching_week(res, dr) <= k and len(days_on(dr)) >= 2,
       f'{n} max {k} in a row: {sorted(days_on(dr), key=ALL.index)} (prev {sorted(dr["w_prev"])})')
ok('Sun' not in days_on(drv(res, 'Rowk00')), 'Rowk00 (worked last Sat, max 1) is off Sunday')
ok('Sun' not in days_on(drv(res, 'Rowk01')), 'Rowk01 (worked last Fri+Sat, max 2) is off Sunday')
ok(not hard_errs(chk), f'no hard-rule errors: {hard_errs(chk)}')
# the verifier catches a violation (simulate one by hand)
dr = drv(res, 'Rowk03')
nxt = next(d for d in ALL[:-1] if d in days_on(dr) and ALL[ALL.index(d) + 1] not in days_on(dr))
dr['bk'].append(ALL[ALL.index(nxt) + 1])          # one day next to a worked day: a run of 2+
chk2 = B.check_invariants(res)
ok(any(e.startswith('PREF-ROW: Rowk03') for e in chk2['errors']), 'verifier: PREF-ROW error on a forced run')
dr['bk'][:] = []
# a run that is ALL last week is not this build's error
res.roster[0]['w_prev'] |= {res.DATEALL['Sun'] - k * ONE for k in range(1, 4)}
ok(not any(e.startswith('PREF-ROW: ' + res.roster[0]['name']) for e in B.check_invariants(res)['errors']),
   'a run entirely in last week is not flagged')

print('A4 keep_together: placement only, days next to each other')
KT = [f'Fairq{i:02d}' for i in range(14)]


def gaps(dr):
    ix = sorted(ALL.index(d) for d in dr['prim'])
    return sum(1 for x, y in zip(ix, ix[1:]) if y - x > 1)


r0, _ = build(synth_week(KT, 6, 'kt0'))
r1, c1 = build(synth_week(KT, 6, 'kt1', keep_together=['Fairq12']))
d0, d1 = drv(r0, 'Fairq12'), drv(r1, 'Fairq12')
ok(gaps(d0) >= 1 and gaps(d1) == 0 and len(d0['prim']) == len(d1['prim']),
   f"Fairq12 {sorted(d0['prim'], key=ALL.index)} -> {sorted(d1['prim'], key=ALL.index)} (same count, together)")
ok(not hard_errs(c1), 'no hard-rule errors')
r2, c2 = build(synth_week(KT, 6, 'pd', prefer_days={'Fairq12': ['Sat']}))
ok('Sat' in drv(r2, 'Fairq12')['prim'] and not hard_errs(c2),
   f"prefer_days: Fairq12 works Sat ({sorted(drv(r2, 'Fairq12')['prim'], key=ALL.index)})")

print('A5 unmatched / ambiguous names: skipped, listed, never an error')
res, chk = build(fixture_cfg(
    strict_names=True,
    extra_unavailable={'Nobody Here': {'Mon': 'timeoff'}, 'Zackary McDonald': {'Mon': 'timeoff'}},
    driver_max_days={'a': 2, 'Daniel Lynch': 3},
    driver_max_row={'Ghost Driver': 1}, keep_together=['Not Real'],
    prefer_days={'Nope Person': ['Sat']}))
um = res.prefs_unmatched
ok(any('Nobody Here' in x for x in um) and any('"a"' in x and 'matches' in x for x in um)
   and any('Ghost Driver' in x for x in um) and any('Not Real' in x for x in um)
   and any('Nope Person' in x for x in um), f'unmatched listed: {um}')
ok(not any('Zackary' in x for x in um), 'an excluded driver is skipped quietly')
ok(total(drv(res, 'Daniel Lynch')) <= 3, 'a matched name in the same key still applies')
buf = io.StringIO()
with contextlib.redirect_stdout(buf):
    B.print_summary(res, chk)
ok("Drivers-page names skipped" in buf.getvalue() and 'Nobody Here' in buf.getvalue(),
   'summary prints the unmatched names')

print('A6 bad shapes fail loudly in load_config')
for bad in ({'extra_unavailable': {'X': ['Mon']}}, {'extra_unavailable': {'X': {'Funday': 'timeoff'}}},
            {'driver_max_days': {'X': -1}}, {'driver_max_row': {'X': 0}},
            {'driver_max_days': {'X': 'two'}}, {'keep_together': 'X'},
            {'prefer_days': {'X': ['Mo']}}):
    try:
        B.load_config(save_cfg(fixture_cfg(**bad), 'bad.json'))
        ok(False, f'rejected {bad}')
    except B.ScheduleConfigError:
        ok(True, f'rejected {bad}')

print('A7 random sweep: the hard caps hold through every pass')
import bk_sweep  # noqa: E402


class _Shim:                      # bk_sweep's case runner, on the CURRENT solver + random prefs
    def __init__(self, rng):
        self.rng = rng

    def load_config(self, p):
        cfg = B.load_config(p)
        rng = self.rng
        names = []
        wb = openpyxl.load_workbook(cfg['avail_file'])
        for r in wb.active.iter_rows(min_row=2, values_only=True):
            names.append(r[0])
        md, mr, eu, kt, pf = {}, {}, {}, [], {}
        for n in names:
            x = rng.random()
            if x < 0.3:
                md[n] = rng.randint(0, 4)
            if rng.random() < 0.2:
                mr[n] = rng.randint(1, 3)
            if rng.random() < 0.15:
                eu[n] = {d: rng.choice(['timeoff', 'days', 'weekend']) for d in rng.sample(ALL, 2)}
            if rng.random() < 0.15:
                kt.append(n)
            if rng.random() < 0.1:
                pf[n] = rng.sample(ALL, 2)
        cfg.update(driver_max_days=md, driver_max_row=mr, extra_unavailable=eu,
                   keep_together=kt, prefer_days=pf)
        return cfg

    def build_schedule(self, cfg):
        return B.build_schedule(cfg)

    def check_invariants(self, res):
        return B.check_invariants(res)


viol = []
for seed in range(1, 4):
    rng = random.Random(seed)
    for k in range(10):
        case = bk_sweep.gen_case(rng, f'pf{seed}x{k}')
        bk_sweep.B = _Shim(random.Random(seed * 100 + k))
        res, chk = bk_sweep.run_case(case, f'pf{seed}x{k}')
        bad = hard_errs(chk)
        for dr in res.roster:
            if dr['max_days'] is not None and total(dr) > max(dr['max_days'], len(dr['extra']) + len(dr['meet'])):
                bad.append(f"{dr['name']} {total(dr)} > max_days {dr['max_days']}")
            if dr['max_row'] is not None and longest_run_touching_week(res, dr) > dr['max_row']:
                bad.append(f"{dr['name']} run > max_row {dr['max_row']}")
            if days_on(dr) & set(dr['unav_why']):
                bad.append(f"{dr['name']} on a Drivers-page day off")
        if bad:
            viol.append((seed, k, bad[:3]))
ok(not viol, f'30 random weeks with random Drivers-page prefs: no hard violations {viol[:2]}')

# ========================================================== B. the engine ====
print('B1 runner: a Drivers-page day off is the typed-confirm bucket')
cfg = fixture_cfg(extra_unavailable={'Colin Drake': {'Mon': 'timeoff', 'Tue': 'days', 'Sat': 'weekend'}},
                  driver_max_days={'Aaron Bell': 1}, driver_max_row={'Bianca Cole': 1},
                  keep_together=['Nobody Atall'])
runner.use_slot('build')
built = json.loads(runner.run(save_cfg(cfg, 'run.json')))
ok(built['ok'], 'build ran')
ok(any('Nobody Atall' in x for x in built['prefs_unmatched']), 'report carries prefs_unmatched')
row = next(r for r in built['drivers'] if r['name'] == 'Colin Drake')
ok(row['unav_why'] == {'Mon': 'timeoff', 'Tue': 'days', 'Sat': 'weekend'}
   and row['cells']['Mon'] == 'Unavailable', 'driver row: unav_why + Unavailable cell')
want = {'Mon': 'asked for this day off (Drivers page)', 'Tue': "can't work Tues (Drivers page)",
        'Sat': 'off weekend (Drivers page)'}
for d, txt in want.items():
    cands = J(runner.candidates, {'day': d, 'role': 'backup'})['candidates']
    cd = next(x for x in cands if x['name'] == 'Colin Drake')
    ok(cd['status'] == 'unavail' and cd['reasons'][0] == txt, f'{d}: {cd["status"]} / {cd["reasons"][:1]}')
r = J(runner.apply_add, {'name': 'Colin Drake', 'day': 'Mon', 'role': 'backup'})
ok(not r['ok'] and r['kind'] == 'needs_confirm', 'needs the typed confirm')
r = J(runner.apply_add, {'name': 'Colin Drake', 'day': 'Mon', 'role': 'backup', 'confirm_unavailable': True})
ok(r['ok'] and any(e.startswith('UNAVAIL violated: Colin Drake Mon') for e in r['check']['overridden'])
   and not any(e.startswith('UNAVAIL') for e in r['check']['errors']), 'confirmed -> an approved override')
ok('asked for this day off (Drivers page)' in r['edits'][-1], f"edit log says why: {r['edits'][-1]}")

print('B2 runner: past max_days is POLICY (allowed, flagged)')
ab = next(x for x in r['drivers'] if x['name'] == 'Aaron Bell')
ok(len(ab['road_days']) + len(ab['backup_days']) <= 1, f"Aaron built within max_days 1: {ab['road_days']} {ab['backup_days']}")
opts = J(runner.add_options, {'name': 'Aaron Bell'})['days']
pick = next((o for o in opts if o['backup']['status'] == 'warn'
             and 'asked for 1 day max (Drivers page)' in o['backup']['reasons']), None)
ok(pick is not None, 'add_options flags a day as policy with "asked for 1 day max (Drivers page)"')
if pick:
    r = J(runner.apply_add, {'name': 'Aaron Bell', 'day': pick['day'], 'role': 'backup'})
    ok(r['ok'], 'allowed without any confirm')
    ok(any(e.startswith('PREF-DAYS: Aaron Bell') for e in r['check']['overridden'])
       and not any(e.startswith('PREF-DAYS') for e in r['check']['errors']),
       f"PREF-DAYS lands in overridden: {r['check']['overridden']}")
    ok('[flagged:' in r['edits'][-1] and 'Drivers page' in r['edits'][-1], f"logged: {r['edits'][-1]}")

print('B3 runner: past max_row is POLICY with its own text')
opts = J(runner.add_options, {'name': 'Bianca Cole'})['days']
hit = [o for o in opts for role in ('road', 'backup')
       if 'asked for no more than 1 day in a row (Drivers page)' in o[role]['reasons']]
ok(bool(hit), 'add_options flags an adjacent day: "asked for no more than 1 day in a row"')
bc = [o for o in hit if o['backup']['status'] == 'warn']
if bc:
    r = J(runner.apply_add, {'name': 'Bianca Cole', 'day': bc[0]['day'], 'role': 'backup'})
    ok(r['ok'] and any(e.startswith('PREF-ROW: Bianca Cole') for e in r['check']['overridden']),
       'PREF-ROW of an approved edit lands in overridden')

print('B4 export_state / load_state keep the new fields')
st = J(runner.export_state)['state']
runner.use_slot('live-prefs')
rep = J(runner.load_state, {'state': st, 'out': os.path.join(tmp, 'live.xlsx')})
ok(rep['ok'], 'reloaded')
row = next(r for r in rep['drivers'] if r['name'] == 'Colin Drake')
ok(row['unav_why'].get('Tue') == 'days', 'unav_why survives the round trip')
ok(next(r for r in rep['drivers'] if r['name'] == 'Aaron Bell')['max_days'] == 1, 'max_days survives')
ok(any(e.startswith('PREF-DAYS: Aaron Bell') for e in rep['check']['overridden']), 'approved override survives')
cd = next(x for x in J(runner.candidates, {'day': 'Tue', 'role': 'backup'})['candidates']
          if x['name'] == 'Colin Drake')
ok(cd['reasons'][0] == "can't work Tues (Drivers page)", 'reason text after reload')
# Route Tracker sync: a Drivers-page day off is reported with the same reason
sat = rep['days'][-1]['date']
tue = next(x['date'] for x in rep['days'] if x['day'] == 'Tue')
sy = J(runner.sync_actual, {'today': sat, 'entries': [
    {'name': 'Colin Drake', 'keys': ['colin|drake'], 'date': tue, 'role': 'road', 'start': '10:45'}]})
sk = [x for x in sy['data']['skipped'] if x['name'] == 'Colin Drake']
ok(sk and sk[0]['reason'] == "can't work Tues (Drivers page)", f'sync_actual skip reason: {sk[:1]}')

print('B5 a week saved BEFORE this change still loads')
old = json.loads(st)
for dr in old['res']['roster']:
    for k in ('unav_why', 'max_days', 'max_row'):
        dr.pop(k, None)
for k in ('prefs_unmatched', 'KEEPT', 'PREFD'):
    old['res'].pop(k, None)
runner.use_slot('live-old')
rep = J(runner.load_state, {'state': json.dumps(old), 'out': os.path.join(tmp, 'old.xlsx')})
ok(rep['ok'] and rep['prefs_unmatched'] == [], f"old state loads: {rep.get('message', '')[:300]}")
cd = next(x for x in J(runner.candidates, {'day': 'Sat', 'role': 'backup'})['candidates']
          if x['name'] == 'Colin Drake')
ok(cd['status'] == 'unavail' and cd['reasons'][0] == 'marked Unavailable that day',
   f"old state: plain Unavailable reason ({cd['status']} {cd['reasons'][:1]})")
r = J(runner.apply_add, {'name': 'Colin Drake', 'day': 'Sat', 'role': 'backup', 'confirm_unavailable': True})
ok(r['ok'] and any(e.startswith('UNAVAIL violated: Colin Drake Sat') for e in r['check']['overridden']),
   f"an edit on the old state works: {r.get('message', '')[:200]}")
ok(J(runner.export_xlsx)['ok'], 'old state: workbook writes')

print()
print('PASS' if not fails else f'FAILED ({len(fails)})')
sys.exit(1 if fails else 0)
